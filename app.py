
import os
import re
import io
import pandas as pd
import streamlit as st
import fitz
from docx import Document
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity


# -------------------- PROJECT CONFIGURATION --------------------

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

MODEL_PATH = os.path.join(
    BASE_DIR,
    "models",
    "all-MiniLM-L6-v2"
)

st.set_page_config(
    page_title="Intelligent ATS",
    page_icon="📄",
    layout="wide"
)


# -------------------- AI MODEL --------------------

@st.cache_resource
def load_model():
    if not os.path.isdir(MODEL_PATH):
        raise FileNotFoundError(
            f"Model folder not found: {MODEL_PATH}"
        )

    return SentenceTransformer(MODEL_PATH)


# -------------------- RESUME TEXT EXTRACTION --------------------

def extract_pdf_text(file_bytes):
    text = ""

    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        for page in pdf:
            text += page.get_text() + "\n"

    return text.strip()


def extract_docx_text(file_bytes):
    document = Document(io.BytesIO(file_bytes))

    paragraphs = [p.text for p in document.paragraphs]

    # Also extract text from tables
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(
                " ".join(cell.text for cell in row.cells)
            )

    return "\n".join(paragraphs).strip()


def extract_resume_text(uploaded_file):
    file_bytes = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    try:
        if filename.endswith(".pdf"):
            return extract_pdf_text(file_bytes)

        elif filename.endswith(".docx"):
            return extract_docx_text(file_bytes)

        else:
            return ""

    except Exception as error:
        st.warning(
            f"Could not read {uploaded_file.name}: {error}"
        )
        return ""


# -------------------- FEATURE ENGINEERING --------------------

SKILLS = [
    "python",
    "sql",
    "mysql",
    "excel",
    "power bi",
    "tableau",
    "machine learning",
    "deep learning",
    "nlp",
    "data analysis",
    "data visualization",
    "statistics",
    "pandas",
    "numpy",
    "scikit-learn",
    "tensorflow",
    "pytorch",
    "java",
    "javascript",
    "html",
    "css",
    "flask",
    "streamlit",
    "communication",
    "project management",
    "leadership",
    "aws",
    "azure",
    "git",
    "github"
]


def normalize_text(text):
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_skills(text):
    normalized = normalize_text(text)
    found_skills = []

    for skill in SKILLS:
        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(skill)
            + r"(?![a-z0-9])"
        )

        if re.search(pattern, normalized):
            found_skills.append(skill)

    return found_skills


def calculate_skill_match(resume_text, job_description):
    resume_skills = set(extract_skills(resume_text))
    job_skills = set(extract_skills(job_description))

    if not job_skills:
        return 0.0, [], []

    matched_skills = sorted(resume_skills & job_skills)
    missing_skills = sorted(job_skills - resume_skills)

    score = (
        len(matched_skills) / len(job_skills)
    ) * 100

    return score, matched_skills, missing_skills


# -------------------- CANDIDATE RANKING --------------------

def rank_resumes(resumes, job_description, model):
    job_description = normalize_text(job_description)

    job_embedding = model.encode(
        [job_description],
        normalize_embeddings=True
    )

    results = []

    for filename, resume_text in resumes:
        clean_resume = normalize_text(resume_text)

        if not clean_resume:
            continue

        skill_score, matched, missing = calculate_skill_match(
            clean_resume,
            job_description
        )

        resume_embedding = model.encode(
            [clean_resume],
            normalize_embeddings=True
        )

        similarity = float(
            cosine_similarity(
                job_embedding,
                resume_embedding
            )[0][0]
        )

        # Convert cosine similarity to a 0–100 score
        semantic_score = max(0.0, min(1.0, similarity)) * 100

        # Combined score: 50% skills + 50% semantic similarity
        final_score = (
            0.50 * skill_score
            + 0.50 * semantic_score
        )

        results.append({
            "Candidate": filename,
            "Skill Match (%)": round(skill_score, 2),
            "Semantic Similarity (%)": round(semantic_score, 2),
            "Final Score (%)": round(final_score, 2),
            "Matched Skills": ", ".join(matched) if matched else "None",
            "Missing Skills": ", ".join(missing) if missing else "None",
            "Resume Text Length": len(clean_resume),
            "Resume Word Count": len(clean_resume.split()),
            "_resume_text": resume_text
        })

    results.sort(
        key=lambda item: item["Final Score (%)"],
        reverse=True
    )

    for rank, result in enumerate(results, start=1):
        result["Rank"] = rank

    return results


# -------------------- USER INTERFACE --------------------

st.title("📄 Intelligent ATS")
st.subheader(
    "AI-Powered Bulk Resume Screening and Candidate Ranking System"
)

st.write(
    "Upload resumes and enter a job description to compare "
    "candidate skills and semantic relevance using NLP."
)

with st.sidebar:
    st.header("About the Project")
    st.write("AI Model: all-MiniLM-L6-v2")
    st.write("Text Features: Skills, word count and text length")
    st.write("Scoring: 50% skill match + 50% semantic similarity")

    st.caption(
        "This is a decision-support tool. Review resumes manually "
        "before making recruitment decisions."
    )


job_description = st.text_area(
    "Enter Job Description",
    height=220,
    placeholder=(
        "Paste the complete job description here, "
        "including responsibilities and required skills..."
    )
)

uploaded_files = st.file_uploader(
    "Upload resumes (PDF or DOCX)",
    type=["pdf", "docx"],
    accept_multiple_files=True
)

if st.button("Screen and Rank Resumes", type="primary"):
    if not job_description.strip():
        st.error("Please enter a job description.")

    elif not uploaded_files:
        st.error("Please upload at least one resume.")

    else:
        with st.spinner("Loading AI model..."):
            try:
                model = load_model()
            except Exception as error:
                st.error(f"Could not load the AI model: {error}")
                st.stop()

        extracted_resumes = []
        extraction_issues = []

        with st.spinner("Extracting resume text..."):
            for uploaded_file in uploaded_files:
                resume_text = extract_resume_text(uploaded_file)

                if resume_text.strip():
                    extracted_resumes.append(
                        (uploaded_file.name, resume_text)
                    )
                else:
                    extraction_issues.append(uploaded_file.name)

        if extraction_issues:
            st.warning(
                "No readable text was extracted from: "
                + ", ".join(extraction_issues)
                + ". Scanned PDFs may need OCR."
            )

        if not extracted_resumes:
            st.error(
                "No readable resume text found. "
                "Please check your uploaded files."
            )
            st.stop()

        with st.spinner("Calculating scores and ranking candidates..."):
            results = rank_resumes(
                extracted_resumes,
                job_description,
                model
            )

        if not results:
            st.error("Could not rank the uploaded resumes.")
            st.stop()

        st.success(
            f"Screening completed for {len(results)} resume(s)."
        )

        # -------------------- SUMMARY --------------------

        best_candidate = results[0]

        col1, col2, col3 = st.columns(3)

        col1.metric("Resumes Processed", len(results))
        col2.metric(
            "Top Candidate Score",
            f"{best_candidate['Final Score (%)']:.2f}%"
        )
        col3.metric("Job Skills Identified", len(
            extract_skills(job_description)
        ))

        # -------------------- RANKING TABLE --------------------

        st.subheader("🏆 Candidate Ranking")

        display_columns = [
            "Rank",
            "Candidate",
            "Skill Match (%)",
            "Semantic Similarity (%)",
            "Final Score (%)",
            "Matched Skills",
            "Missing Skills"
        ]

        display_df = pd.DataFrame(results)[display_columns]

        st.dataframe(
            display_df,
            use_container_width=True,
            hide_index=True
        )

        # -------------------- CANDIDATE DETAILS --------------------

        st.subheader("🔍 Candidate Details")

        for result in results:
            with st.expander(
                f"Rank {result['Rank']} — {result['Candidate']} "
                f"({result['Final Score (%)']:.2f}%)"
            ):
                st.write(
                    f"**Skill Match:** {result['Skill Match (%)']}%"
                )
                st.write(
                    f"**Semantic Similarity:** "
                    f"{result['Semantic Similarity (%)']}%"
                )
                st.write(
                    f"**Resume Word Count:** "
                    f"{result['Resume Word Count']}"
                )
                st.write(
                    f"**Resume Text Length:** "
                    f"{result['Resume Text Length']} characters"
                )
                st.write(
                    f"**Matched Skills:** {result['Matched Skills']}"
                )
                st.write(
                    f"**Missing Skills:** {result['Missing Skills']}"
                )

                with st.expander("View Extracted Resume Text"):
                    st.text(result["_resume_text"])

        # -------------------- CSV REPORT --------------------

        st.subheader("📥 Download Report")

        report_df = pd.DataFrame(results)[display_columns]

        csv_data = report_df.to_csv(
            index=False
        ).encode("utf-8-sig")

        st.download_button(
            label="Download Candidate Ranking CSV",
            data=csv_data,
            file_name="ATS_Candidate_Ranking.csv",
            mime="text/csv"
        )

        st.info(
            "Scores indicate similarity to the supplied job description. "
            "They are not a guarantee of candidate quality or job performance."
        )
